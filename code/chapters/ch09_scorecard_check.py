# Chapter 9: print the region scorecard after LibreOffice has recalculated it
# openpyxl reads the values Excel or LibreOffice saved, so open and save the
# workbook in Excel first, or recalculate it with LibreOffice:
#   soffice --headless --convert-to xlsx --outdir recalc templates/region_scorecard.xlsx
import sys
from openpyxl import load_workbook

path = sys.argv[1] if len(sys.argv) > 1 else "recalc/region_scorecard.xlsx"
ws = load_workbook(path, data_only=True).active
rows = list(ws.iter_rows(values_only=True))
print(f"{'Region':<7}{'FY26 $k':>9}{'Target $k':>11}{'FY25 $k':>9}"
      f"{'Var $k':>8}{'Var %':>8}{'vs LY %':>9}  Flag")
for r in rows[1:]:
    if r[0] is None:  # stop at the first blank row (notes follow)
        break
    print(f"{r[0]:<7}{r[1]/1000:>9,.0f}{r[2]/1000:>11,.0f}{r[3]/1000:>9,.0f}"
          f"{r[4]/1000:>8,.1f}{r[5]:>+8.1%}{r[6]:>+9.1%}  {r[7] or ''}".rstrip())
